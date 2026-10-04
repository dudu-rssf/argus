# -*- coding: utf-8 -*-
"""Gera docs/catalogo/Argus_Catalogo_de_Dados.xlsx a partir dos CSVs (fonte da verdade).

Uso (a partir de pipeline/): uv run --with openpyxl python -m argus_pipeline.catalog_xlsx
"""
import csv
from pathlib import Path

CAT = Path(__file__).resolve().parents[2] / "docs" / "catalogo"


def _ler(nome):
    with open(CAT / f"{nome}.csv", encoding="utf-8", newline="") as f:
        return list(csv.reader(f))[1:]


_BR, _US = _ler("brasil"), _ler("eua")
# formato interno: (aba, bloco, indicador, fonte, codigo, freq, unidade, tipo, fase, status, obs) + transformações
BR = [tuple(r[1:9] + r[10:13]) for r in _BR]
US = [tuple(r[1:9] + r[10:13]) for r in _US]
TRANSF_LINHA = {r[0]: r[9] for r in _BR + _US}
CENTRAL = [tuple(r) for r in _ler("central")]
FONTES = [tuple(r) for r in _ler("fontes")]
TRANSF = {}
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

NAVY, GOLD, LIGHT = "0B1F3A", "C9A227", "F3F5F8"
F = "Arial"
hfont = Font(name=F, bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", fgColor=NAVY)
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=16, bold=True, color=NAVY)
sub = Font(name=F, size=10, italic=True, color="5A6B82")
thin = Side(style="thin", color="D5DBE3")
border = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
STATUS = ["Verificado", "Confirmar", "A mapear", "Sem API", "Pago", "Derivado", "Curado"]
FASES = ["MVP", "EUA", "v2", "v3"]
STATUS_COLOR = {"Verificado": "D8EFDF", "Confirmar": "FFF4D1", "A mapear": "FDE2CF",
                "Sem API": "ECE6F5", "Pago": "F8D7DA", "Derivado": "DCE8F7", "Curado": "EEEEEE"}

wb = Workbook()

def header_block(ws, t, s, ncols):
    ws["A1"] = t; ws["A1"].font = title
    ws["A2"] = s; ws["A2"].font = sub
    ws.row_dimensions[1].height = 26
    for c in range(1, ncols + 1):
        ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor=GOLD)
    ws.row_dimensions[3].height = 3

def table(ws, headers, rows, widths, start=5):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=start, column=i, value=h)
        c.font, c.fill = hfont, hfill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[start].height = 28
    for r, row in enumerate(rows, start + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.alignment, c.border = body, wrap, border
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    end = start + len(rows)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{end}"
    ws.freeze_panes = ws.cell(row=start + 1, column=4)
    return end

def catalog(ws, prefix, data, t, s):
    headers = ["ID", "Aba", "Bloco", "Indicador", "Fonte", "Código / endpoint", "Frequência",
               "Unidade", "Tipo", "Transformações válidas", "Fase", "Status", "Observação"]
    rows = []
    for n, (aba, bloco, ind, fonte, cod, freq, un, tipo, fase, st, obs) in enumerate(data, 1):
        rid = f"{prefix}-{n:03d}"
        rows.append([rid, aba, bloco, ind, fonte, cod, freq, un, tipo,
                     TRANSF_LINHA.get(rid, ""), fase, st, obs])
    header_block(ws, t, s, len(headers))
    end = table(ws, headers, rows, [9, 18, 14, 42, 20, 30, 14, 11, 12, 34, 7, 12, 40])
    for r in range(6, end + 1):
        ws.cell(row=r, column=1).font = Font(name=F, size=9, color="5A6B82")
        ws.cell(row=r, column=4).font = bold
    dv_f = DataValidation(type="list", formula1='"' + ",".join(FASES) + '"', allow_blank=True)
    dv_s = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv_f); ws.add_data_validation(dv_s)
    dv_f.add(f"K6:K{end+200}"); dv_s.add(f"L6:L{end+200}")
    for st, col in STATUS_COLOR.items():
        ws.conditional_formatting.add(f"L6:L{end+200}", FormulaRule(
            formula=[f'$L6="{st}"'], fill=PatternFill("solid", fgColor=col)))
    ws.conditional_formatting.add(f"K6:K{end+200}", FormulaRule(
        formula=['$K6="MVP"'], font=Font(name=F, bold=True, color=NAVY)))
    return end

# ---------- Leia-me ----------
ws = wb.active; ws.title = "Leia-me"
header_block(ws, "Argus — Catálogo de Dados", "Fonte única das séries de cada aba. Vira o arquivo de configuração do pipeline na Fase 1.", 3)
rows = [
 ("Como usar", ""),
 ("Brasil / EUA", "Uma linha por indicador. Filtre por Aba, Fase ou Status. As colunas Fase e Status têm lista suspensa e podem ser editadas."),
 ("Central", "Blocos da aba Central (Command Center) e do agregador de notícias."),
 ("Fontes", "Registro de cada fonte: acesso, autenticação, limites e risco técnico."),
 ("Resumo", "Contagens por aba, fase e status, calculadas por fórmula."),
 ("", ""),
 ("Status", ""),
 ("Verificado", "Código conferido nesta sessão contra o portal oficial (título ou último valor)."),
 ("Confirmar", "Código conhecido, ainda não conferido. A validação automática da Fase 0 confirma título e última observação."),
 ("A mapear", "Fonte definida, código ainda não localizado."),
 ("Sem API", "Dado gratuito, mas só por download manual ou PDF. Entra como carga manual ou fica para v2."),
 ("Pago", "Exige licença. A coluna Observação indica o substituto gratuito quando existe."),
 ("Derivado", "Calculado pelo Argus a partir de outras séries. A fórmula está em Código / endpoint."),
 ("Curado", "Tabela mantida manualmente (decisões do Copom, metas, pesquisas)."),
 ("", ""),
 ("Fase", ""),
 ("MVP", "Brasil + Central. Fontes oficiais com API, sem modelagem própria."),
 ("EUA", "Logo após o MVP: mesmo template de abas, quase tudo via FRED."),
 ("v2", "Fontes difíceis (B3, ANBIMA, pesquisas, texto) ou modelos próprios."),
 ("v3", "Projetos de maior risco: NLP de discursos, fluxo B3, balanços CVM."),
 ("", ""),
 ("Tipo", "Define quais transformações o motor de gráfico oferece. Ex.: taxa nunca vira acumulado 12m; variação mensal é encadeada, não somada."),
 ("Data", "Planilha gerada a partir dos CSVs em docs/catalogo/. Edite os CSVs e rode o gerador; resultado da validação em validacao.md."),
]
for i, (a, b) in enumerate(rows, 5):
    ws.cell(row=i, column=1, value=a).font = bold if b == "" or a in STATUS + FASES + ["Tipo", "Data", "Brasil / EUA", "Central", "Fontes", "Resumo"] else body
    ws.cell(row=i, column=2, value=b).font = body
    ws.cell(row=i, column=2).alignment = wrap
    if a in STATUS_COLOR:
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=STATUS_COLOR[a])
    if b == "" and a:
        ws.cell(row=i, column=1).font = Font(name=F, size=11, bold=True, color=NAVY)
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 110

# ---------- Resumo ----------
wr = wb.create_sheet("Resumo")
# ---------- Catálogos ----------
wbr = wb.create_sheet("Brasil")
br_end = catalog(wbr, "BR", BR, "Brasil — catálogo de séries", "Subabas da aba Brasil. Foco do MVP.")
wus = wb.create_sheet("EUA")
us_end = catalog(wus, "US", US, "EUA — catálogo de séries", "Mapeamento da planilha original de dados US para fonte e código.")

# ---------- Central ----------
wc = wb.create_sheet("Central")
header_block(wc, "Aba Central — Command Center e notícias", "O que aparece na primeira tela e de onde vem.", 6)
table(wc, ["Bloco", "Componente", "Conteúdo", "Fonte", "Fase", "Observação"], CENTRAL, [18, 20, 60, 36, 8, 46])
dv = DataValidation(type="list", formula1='"' + ",".join(FASES) + '"', allow_blank=True)
wc.add_data_validation(dv); dv.add("E6:E100")

# ---------- Fontes ----------
wf = wb.create_sheet("Fontes")
header_block(wf, "Registro de fontes", "Cada fonte vira um adapter isolado no pipeline.", 6)
table(wf, ["Fonte", "Acesso", "Autenticação", "Limite / observação", "Risco técnico", "Custo"], FONTES, [32, 30, 22, 56, 13, 14])
for r in range(6, 6 + len(FONTES)):
    v = wf.cell(row=r, column=5).value
    col = {"Baixo": "D8EFDF", "Médio": "FFF4D1", "Alto": "F8D7DA"}.get(v)
    if col: wf.cell(row=r, column=5).fill = PatternFill("solid", fgColor=col)

# ---------- Resumo (fórmulas) ----------
header_block(wr, "Resumo do catálogo", "Contagens calculadas a partir das abas Brasil e EUA.", 9)
def summary(ws, top, sheet, end, label):
    abas = []
    data = BR if sheet == "Brasil" else US
    for d in data:
        if d[0] not in abas: abas.append(d[0])
    ws.cell(row=top, column=1, value=label).font = Font(name=F, size=12, bold=True, color=NAVY)
    hdr = ["Aba"] + FASES + ["Total"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=top + 1, column=i, value=h); c.font, c.fill = hfont, hfill
    rng_a = f"{sheet}!$B$6:$B${end}"; rng_f = f"{sheet}!$K$6:$K${end}"
    for j, a in enumerate(abas):
        r = top + 2 + j
        ws.cell(row=r, column=1, value=a).font = body
        for k, f in enumerate(FASES):
            ws.cell(row=r, column=2 + k, value=f'=COUNTIFS({rng_a},$A{r},{rng_f},"{f}")').font = body
        ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})").font = bold
    tr = top + 2 + len(abas)
    ws.cell(row=tr, column=1, value="Total").font = bold
    for k in range(5):
        col = get_column_letter(2 + k)
        ws.cell(row=tr, column=2 + k, value=f"=SUM({col}{top+2}:{col}{tr-1})").font = bold
    # status block
    ws.cell(row=top + 1, column=8, value="Status").font = hfont
    ws.cell(row=top + 1, column=8).fill = hfill
    ws.cell(row=top + 1, column=9, value="Qtd").font = hfont
    ws.cell(row=top + 1, column=9).fill = hfill
    for j, st in enumerate(STATUS):
        r = top + 2 + j
        ws.cell(row=r, column=8, value=st).font = body
        ws.cell(row=r, column=8).fill = PatternFill("solid", fgColor=STATUS_COLOR[st])
        ws.cell(row=r, column=9, value=f'=COUNTIF({sheet}!$L$6:$L${end},H{r})').font = body
    return max(tr, top + 2 + len(STATUS)) + 2

nxt = summary(wr, 5, "Brasil", br_end, "Brasil")
summary(wr, nxt, "EUA", us_end, "EUA")
wr.column_dimensions["A"].width = 24
for c in "BCDEF": wr.column_dimensions[c].width = 9
wr.column_dimensions["G"].width = 4; wr.column_dimensions["H"].width = 14; wr.column_dimensions["I"].width = 8

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.move_sheet("Resumo", offset=0)
out = str(CAT / "Argus_Catalogo_de_Dados.xlsx")
wb.save(out); print(out, len(BR), len(US))
